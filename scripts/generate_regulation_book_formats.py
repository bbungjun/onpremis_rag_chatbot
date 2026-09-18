"""한 합성 규정집을 DOCX와 선택 가능한 텍스트 PDF로 만든다."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
from datetime import UTC, datetime
from pathlib import Path

from docx import Document
from docx.shared import Mm, Pt, RGBColor
from pypdf import PdfReader
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app.chunking import chunk_text
from app.document_reader import read_document
from scripts.generate_hwp_corpus import make_regulation_book

TITLE = "가상회사 합성 사내 규정집"
SUBTITLE = "파싱·검색 실험용 문서이며 실제 사내 규정이 아닙니다."
_ARTICLE = re.compile(r"^\*\*(제\d+조(?:의\d+)?(?:\s*\([^)]*\))?)\*\*$")


def generate_formats(count: int, output: Path, *, font_path: Path | None = None) -> dict:
    """정책 한 원문을 두 파일로 쓰고 두 형식의 조·항/질문 매핑을 검증한다."""
    book = make_regulation_book(count)
    output = Path(output)
    stem = f"synthetic-regulations-{count}"
    docx_path = output / f"{stem}.docx"
    pdf_path = output / f"{stem}.pdf"
    manifest_path = output / "manifest.json"
    for path in (docx_path, pdf_path, manifest_path):
        if path.exists():
            raise FileExistsError(f"기존 생성물을 덮어쓰지 않습니다: {path}")

    output.mkdir(parents=True, exist_ok=True)
    blocks = book.markdown.strip().split("\n\n")
    _write_docx(blocks, docx_path)
    _write_pdf(blocks, pdf_path, font_path)

    expected = chunk_text(book.markdown)
    parent_count = sum(chunk["type"] == "parent" for chunk in expected)
    child_count = sum(chunk["type"] == "child" for chunk in expected)
    if (parent_count, child_count) != (8 * count, 24 * count):
        raise ValueError("원본 합성 규정의 조·항 수가 예상과 다릅니다")
    if len({chunk["id"] for chunk in expected}) != len(expected):
        raise ValueError("원본 합성 규정의 청크 ID가 중복됐습니다")
    for path in (docx_path, pdf_path):
        _verify_roundtrip(path, expected, book.questions)

    result = {
        "policies": count,
        "parent_chunks": parent_count,
        "child_chunks": child_count,
        "qa_matches": len(book.questions),
        "pdf_pages": len(PdfReader(pdf_path).pages),
        "source_markdown_sha256_lf": hashlib.sha256(book.markdown.encode("utf-8")).hexdigest(),
        "font": font_path.name if font_path else "HYSMyeongJo-Medium",
        "files": {
            path.suffix.lstrip("."): {
                "path": path.name,
                "bytes": path.stat().st_size,
                "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
            }
            for path in (docx_path, pdf_path)
        },
    }
    manifest_path.write_text(
        json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    return result


def _write_docx(blocks: list[str], output: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = section.bottom_margin = Mm(18)
    section.left_margin = section.right_margin = Mm(20)
    for name, size, after in (
        ("Title", 18, 12),
        ("Heading 1", 15, 9),
        ("Heading 2", 12.5, 7),
        ("Heading 3", 11, 6),
        ("Heading 4", 10.5, 4),
        ("Normal", 10, 3),
    ):
        style = document.styles[name]
        style.font.name = "Malgun Gothic"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_after = Pt(after)
        if name.startswith("Heading"):
            style.paragraph_format.keep_with_next = True

    document.add_paragraph(TITLE, style="Title")
    document.add_paragraph(SUBTITLE)
    part_count = 0
    for block in blocks:
        if block.startswith("# "):
            part_count += 1
            paragraph = document.add_paragraph(block[2:], style="Heading 1")
            paragraph.paragraph_format.page_break_before = part_count > 1
        elif block.startswith("## "):
            document.add_paragraph(block[3:], style="Heading 2")
        elif block.startswith("### "):
            document.add_paragraph(block[4:], style="Heading 3")
        elif article := _ARTICLE.fullmatch(block):
            document.add_paragraph(article.group(1), style="Heading 4")
        else:
            document.add_paragraph(block)

    document.core_properties.title = TITLE
    document.core_properties.subject = SUBTITLE
    document.core_properties.author = "가상회사 합성 데이터"
    document.core_properties.created = datetime(2026, 9, 19, tzinfo=UTC)
    document.save(output)


def _write_pdf(blocks: list[str], output: Path, font_path: Path | None) -> None:
    if font_path is None:
        font_name = "HYSMyeongJo-Medium"
        pdfmetrics.registerFont(UnicodeCIDFont(font_name))
    else:
        font_name = "RegulationBookFont"
        pdfmetrics.registerFont(TTFont(font_name, str(font_path)))

    pdf = canvas.Canvas(str(output), pagesize=A4, pageCompression=1, invariant=1)
    pdf.setTitle(TITLE)
    pdf.setSubject(SUBTITLE)
    pdf.setAuthor("가상회사 합성 데이터")
    page_width, page_height = A4
    left = 20 * mm
    right = 20 * mm
    top = 18 * mm
    bottom = 18 * mm
    usable_width = page_width - left - right
    y = page_height - top
    page_count = 1
    part_count = 0

    def new_page() -> None:
        nonlocal y, page_count
        pdf.showPage()
        page_count += 1
        y = page_height - top

    def draw_block(
        text: str, size: float, leading: float, after: float, heading: bool = False
    ) -> None:
        nonlocal y
        lines = _wrap(text, font_name, size, usable_width)
        required = len(lines) * leading + after + (leading if heading else 0)
        if y - required < bottom:
            new_page()
        pdf.setFont(font_name, size)
        for line in lines:
            if y - leading < bottom:
                new_page()
                pdf.setFont(font_name, size)
            pdf.drawString(left, y, line)
            y -= leading
        y -= after

    draw_block(TITLE, 17, 24, 10, heading=True)
    draw_block(SUBTITLE, 9, 14, 16)
    for block in blocks:
        if block.startswith("# "):
            part_count += 1
            if part_count > 1:
                new_page()
            draw_block(block[2:], 14, 21, 11, heading=True)
        elif block.startswith("## "):
            draw_block(block[3:], 12, 18, 8, heading=True)
        elif block.startswith("### "):
            draw_block(block[4:], 10.5, 16, 7, heading=True)
        elif article := _ARTICLE.fullmatch(block):
            draw_block(article.group(1), 10, 15, 5, heading=True)
        else:
            draw_block(block, 9.5, 14, 5)
    pdf.save()
    if page_count <= 0:
        raise ValueError("PDF 페이지를 생성하지 못했습니다")


def _wrap(text: str, font_name: str, size: float, width: float) -> list[str]:
    lines: list[str] = []
    current = ""
    for word in text.split():
        candidate = f"{current} {word}" if current else word
        if pdfmetrics.stringWidth(candidate, font_name, size) <= width:
            current = candidate
        elif current:
            lines.append(current)
            current = word
        else:
            current = word
        if pdfmetrics.stringWidth(current, font_name, size) > width:
            raise ValueError(f"PDF 페이지 폭을 초과하는 단어: {word[:30]}")
    if current:
        lines.append(current)
    return lines


def _verify_roundtrip(path: Path, expected: list[dict], questions: list[dict]) -> None:
    actual = chunk_text(read_document(path))
    expected_children = {
        chunk["id"]: _compact(chunk["text"]) for chunk in expected if chunk["type"] == "child"
    }
    actual_children = {
        chunk["id"]: _compact(chunk["text"]) for chunk in actual if chunk["type"] == "child"
    }
    if actual_children != expected_children:
        missing = sorted(expected_children.keys() - actual_children.keys())[:3]
        changed = [
            key
            for key in expected_children.keys() & actual_children.keys()
            if expected_children[key] != actual_children[key]
        ][:3]
        raise ValueError(f"{path.suffix} 항 텍스트 불일치: missing={missing}, changed={changed}")
    expected_parent_ids = {chunk["id"] for chunk in expected if chunk["type"] == "parent"}
    parents = {chunk["id"]: chunk for chunk in actual if chunk["type"] == "parent"}
    if (
        set(parents) != expected_parent_ids
        or len(actual) != len(expected)
        or len({chunk["id"] for chunk in actual}) != len(actual)
    ):
        raise ValueError(f"{path.suffix} 조·항 구조가 원본과 다릅니다")
    for question in questions:
        parent = parents[question["expected_chunk_id"]]
        answer = question["expected_answer"].split(". (")[0]
        if _compact(answer) not in _compact(parent["text"]):
            raise ValueError(f"{path.suffix} 답 문구 누락: {question['expected_chunk_id']}")


def _compact(value: str) -> str:
    return re.sub(r"\s+", "", value)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="Make one synthetic regulation book as DOCX and PDF"
    )
    parser.add_argument("--count", type=int, default=1000)
    parser.add_argument("--output", type=Path, default=Path("output/regulation-book-1000"))
    parser.add_argument("--font-path", type=Path, help="Embedded TrueType font for Korean PDF text")
    args = parser.parse_args(argv)
    result = generate_formats(args.count, args.output, font_path=args.font_path)
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
