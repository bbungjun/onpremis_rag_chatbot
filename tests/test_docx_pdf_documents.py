from base64 import b64decode
from io import BytesIO
from pathlib import Path
from shutil import copyfile
from types import SimpleNamespace

import pytest

from app.chunking import chunk_text
from app.document_reader import read_document
from scripts import ingest_md

FIXTURES = Path(__file__).parent / "fixtures"


def _child(chunks, article_id):
    return next(
        chunk["text"]
        for chunk in chunks
        if chunk["type"] == "child" and chunk["parent_id"] == article_id
    )


def test_discovery_includes_docx_and_pdf_case_insensitively(tmp_path):
    (tmp_path / "leave.DOCX").write_bytes(b"fixture")
    (tmp_path / "travel.PDF").write_bytes(b"fixture")

    assert [path.name for path in ingest_md.discover_source_files(tmp_path)] == [
        "leave.DOCX",
        "travel.PDF",
    ]


def test_docx_preserves_article_order_and_table_relationship():
    text = read_document(FIXTURES / "synthetic-policy.docx")
    chunks = chunk_text(text)

    assert [chunk["id"] for chunk in chunks if chunk["type"] == "parent"] == [
        "jo-1",
        "jo-2",
    ]
    assert "3영업일 전까지" in _child(chunks, "jo-1")
    assert "구분: 연차" in _child(chunks, "jo-1")
    assert "신청 기한: 3영업일 전" in _child(chunks, "jo-1")
    assert "5영업일 이내" in _child(chunks, "jo-2")


def test_text_pdf_preserves_article_content():
    text = read_document(FIXTURES / "synthetic-policy.pdf")
    chunks = chunk_text(text)

    assert [chunk["id"] for chunk in chunks if chunk["type"] == "parent"] == [
        "jo-1",
        "jo-2",
    ]
    assert "3영업일 전까지" in _child(chunks, "jo-1")
    assert "5영업일 이내" in _child(chunks, "jo-2")


def test_pdf_without_text_fails_with_ocr_guidance(tmp_path):
    from pypdf import PdfWriter

    source = tmp_path / "scan.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    with source.open("wb") as output:
        writer.write(output)

    with pytest.raises(ValueError, match="OCR"):
        read_document(source)


def test_docx_automatic_numbering_is_rejected_until_preserved(tmp_path):
    from docx import Document

    source = tmp_path / "numbered.docx"
    document = Document()
    document.add_paragraph("제1조 (신청 절차)")
    document.add_paragraph("① 승인 절차는 다음과 같다.")
    document.add_paragraph("팀장에게 신청한다.", style="List Number")
    document.save(source)

    with pytest.raises(ValueError, match="자동 번호"):
        read_document(source)


def test_docx_merged_table_is_rejected_until_preserved(tmp_path):
    from docx import Document

    source = tmp_path / "merged.docx"
    document = Document()
    document.add_paragraph("제1조 (증빙)")
    document.add_paragraph("① 아래 표를 따른다.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "병합 헤더"
    table.cell(1, 0).text = "교통비"
    table.cell(1, 1).text = "영수증"
    document.save(source)

    with pytest.raises(ValueError, match="병합"):
        read_document(source)


def test_docx_inline_image_is_rejected_until_ocr_exists(tmp_path):
    from docx import Document

    source = tmp_path / "image.docx"
    document = Document()
    document.add_paragraph("제1조 (안내)")
    document.add_paragraph("① 다음 그림을 따른다.")
    png = b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+/aL8AAAAASUVORK5CYII="
    )
    document.add_picture(BytesIO(png))
    document.save(source)

    with pytest.raises(ValueError, match="이미지"):
        read_document(source)


def test_same_stem_docx_and_pdf_fail_before_reset(tmp_path, monkeypatch):
    (tmp_path / "leave.docx").write_bytes(b"fixture")
    (tmp_path / "leave.pdf").write_bytes(b"fixture")

    def unreadable(_path):
        raise AssertionError("Document was read before duplicate detection")

    monkeypatch.setattr(ingest_md, "read_document", unreadable)
    deleted = []
    monkeypatch.setattr(ingest_md, "delete_collection_if_exists", lambda *_: deleted.append(True))
    settings = SimpleNamespace(qdrant_url="http://qdrant.test", qdrant_collection="chunks")

    with pytest.raises(ValueError, match="Ambiguous source variants"):
        ingest_md.ingest_directory(tmp_path, settings=settings, reset=True)

    assert deleted == []


def test_docx_and_pdf_flow_to_qdrant_payloads(tmp_path, monkeypatch):
    docs = tmp_path / "docs"
    docs.mkdir()
    copyfile(FIXTURES / "synthetic-policy.docx", docs / "leave.docx")
    copyfile(FIXTURES / "synthetic-policy.pdf", docs / "travel.pdf")
    monkeypatch.setattr(ingest_md, "embed_text", lambda *_: [0.1, 0.2])
    monkeypatch.setattr(ingest_md, "text_to_sparse", lambda *_: {"indices": [1], "values": [1.0]})
    monkeypatch.setattr(ingest_md, "ensure_collection", lambda *_, **__: None)
    uploaded = []
    monkeypatch.setattr(
        ingest_md,
        "upsert_chunk_vectors",
        lambda _url, _collection, points: uploaded.extend(points),
    )
    settings = SimpleNamespace(
        ollama_base_url="http://ollama.test",
        embedding_model="bge-m3",
        qdrant_url="http://qdrant.test",
        qdrant_collection="chunks",
    )

    result = ingest_md.ingest_directory(docs, settings=settings, batch_size=2)

    assert (result.documents_indexed, result.chunks_created, result.vectors_inserted) == (
        2,
        8,
        4,
    )
    assert len(uploaded) == 4
    assert {point["payload"]["title"] for point in uploaded} == {
        "leave.docx",
        "travel.pdf",
    }
